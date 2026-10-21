import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# Usage: python3 tools/build_resume.py <content.py> <output dir>
# content.py defines NAME, CONTACT, BASENAME, and SECTIONS, plus an optional HEADLINE.
CONTENT, OUT = sys.argv[1], sys.argv[2]
cfg = {}
exec(open(CONTENT).read(), cfg)
NAME, CONTACT, BASENAME, SECTIONS = cfg["NAME"], cfg["CONTACT"], cfg["BASENAME"], cfg["SECTIONS"]
HEADLINE = cfg.get("HEADLINE")

# ---------- Markdown ----------
md = [f"# {NAME}", ""] + ([f"**{HEADLINE}**", ""] if HEADLINE else []) + [CONTACT, ""]
for title, entries in SECTIONS:
    md += [f"## {title}", ""]
    for e in entries:
        if isinstance(e, dict):
            md += [f"**{e['org']}** | {e['right']}  ", f"*{e['title']}*", ""]
            md += [f"- {b}" for b in e["bullets"]] + [""]
        else:
            _, lead, rest = e
            md.append(f"- **{lead}:** {rest}")
    if not isinstance(entries[-1], dict):
        md.append("")
open(f"{OUT}/resume.md", "w").write("\n".join(md).rstrip() + "\n")

# ---------- DOCX ----------
FONT = "Calibri"
BODY = 10
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Inches(0.6))
sec.top_margin = sec.bottom_margin = Inches(0.5)
RIGHT_TAB = sec.page_width - sec.left_margin - sec.right_margin

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(BODY)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0

def para(space_before=0, space_after=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    return p

def run(p, text, bold=False, italic=False, size=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    return r

def bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "000000")):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b)
    pPr.append(bdr)

def bullet(p_text_parts):
    p = para()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.2)
    pf.first_line_indent = Inches(-0.13)
    run(p, "• ")
    for text, bold in p_text_parts:
        run(p, text, bold=bold)
    return p

p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
run(p, NAME, bold=True, size=16)
if HEADLINE:
    p = para(align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, HEADLINE, size=10.5)
p = para(space_after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
run(p, CONTACT, size=9.5)

for title, entries in SECTIONS:
    p = para(space_before=6, space_after=2)
    run(p, title.upper(), bold=True, size=10.5)
    bottom_border(p)
    for i, e in enumerate(entries):
        if isinstance(e, dict):
            p = para(space_before=3 if i else 0)
            p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
            run(p, e["org"], bold=True)
            run(p, "\t" + e["right"])
            p = para()
            run(p, e["title"], italic=True)
            for b in e["bullets"]:
                bullet([(b, False)])
        else:
            _, lead, rest = e
            bullet([(lead + ": ", True), (rest, False)])

doc.save(f"{OUT}/{BASENAME}.docx")
print("ok")

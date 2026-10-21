"""Build the L3Harris Quality Engineer Intern cover letter as DOCX (same header as the resume)."""
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(1)
sec.top_margin = sec.bottom_margin = Inches(0.8)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.0

props = doc.core_properties
props.title = "Charles Morrisey Cover Letter"
props.author = "Charles Morrisey"


def para(text="", size=None, align=None, after=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        if size:
            r.font.size = size
    return p


para("CHARLES MORRISEY", size=Pt(20), align=WD_ALIGN_PARAGRAPH.CENTER)
para("San Luis Obispo, CA | (626) 705-1471 | cmorrise@calpoly.edu | linkedin.com/in/charlesmorrisey",
     size=Pt(10), align=WD_ALIGN_PARAGRAPH.CENTER, after=18)

para("October 21, 2026", after=12)
para("Hiring Manager")
para("L3Harris Technologies")
para("Rochester, NY", after=12)
para("Re: Quality Engineer Intern (Job ID 41262)", after=12)
para("Dear Hiring Manager,", after=10)

BODY = [
    "I am applying for the Quality Engineer Intern role in Rochester. I study Biomedical Engineering at "
    "Cal Poly, I am a former West Point cadet, and I hold an inactive DoD Secret clearance that stays "
    "eligible for reactivation through January 2028. I have also spent three summers doing the core of "
    "what your posting describes: inspecting work, documenting nonconformances, and making sure defects "
    "get fixed before the customer signs off.",

    "As a Quality Control Supervisor at Quality Windows, I inspected more than 30 residential and "
    "commercial properties each season against client requirements. When a job fell short, I wrote down "
    "the problem, directed the rework, and held sign-off until the rework was done. I also trained three "
    "employees on safety procedures and equipment use. Those summers taught me that a defect caught on "
    "site costs far less than one the customer finds.",

    "I have kept building on that work. At Coastal Shine, the exterior cleaning business I started in "
    "March, I wrote the standard operating procedures for job setup, cleaning, and final inspection, and "
    "I prepare the cost estimate for every job. At MeldMD, I cleaned and validated a 3,050-record "
    "database and verified contacts for 98% of it, which takes the same care as an inspection log. On "
    "Cal Poly's EMPOWER design team, I turn user needs into written design requirements and test "
    "3D-printed prototypes. I also build part models, assemblies, and engineering drawings in SolidWorks.",

    "I want to learn how a Material Review Board, root cause analysis, and a Zero Defect program run on "
    "a real production line, and L3Harris would let me do that on hardware that supports national "
    "security. I am available from mid-June 2027, after Cal Poly's spring quarter ends, and I am ready "
    "to relocate to Rochester. Thank you for your time. I would welcome the chance to talk about how I "
    "can support your Quality team.",
]
for text in BODY:
    p = para(text, after=10)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

para("Sincerely,", after=24)
para("Charles Morrisey")

out = sys.argv[1] if len(sys.argv) > 1 else "Charles_Morrisey_L3Harris_Quality_Engineer_Intern_Cover_Letter.docx"
doc.save(out)
print("wrote", out)
